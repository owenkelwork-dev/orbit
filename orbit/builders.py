"""Turn FORGE's product specs into real files: .xlsx spreadsheets and printable PDFs."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def check_product_spec(spec: Any) -> list[str]:
    if not isinstance(spec, dict):
        return ["product must be an object"]
    kind = spec.get("type")
    if kind == "spreadsheet":
        sheets = spec.get("sheets")
        if not isinstance(sheets, list) or not sheets:
            return ["spreadsheet product needs a non-empty sheets list"]
        errors = []
        for s in sheets:
            if not isinstance(s.get("columns"), list) or not s["columns"]:
                errors.append(f"sheet {s.get('name')!r} needs a non-empty columns list")
            if not isinstance(s.get("rows", []), list):
                errors.append(f"sheet {s.get('name')!r} rows must be a list of lists")
        return errors
    if kind == "printable":
        pages = spec.get("pages")
        if not isinstance(pages, list) or not pages:
            return ["printable product needs a non-empty pages list"]
        return [f"page {i + 1} needs a title" for i, p in enumerate(pages) if not p.get("title")]
    if kind == "other":
        return []
    return ['product.type must be "spreadsheet", "printable", or "other"']


def build_spreadsheet(spec: dict, path: Path) -> Path:
    """Spec: {sheets: [{name, columns, rows, column_widths?, notes?}]}. Cells starting with '=' are formulas."""
    wb = Workbook()
    wb.remove(wb.active)
    accent = PatternFill("solid", fgColor=spec.get("header_color", "2F5D62").lstrip("#"))
    for sheet in spec["sheets"]:
        ws = wb.create_sheet(str(sheet.get("name", "Sheet"))[:31])
        ws.append(list(sheet["columns"]))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = accent
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.get("rows", []):
            ws.append(list(row) if isinstance(row, list) else [row])
        widths = sheet.get("column_widths") or []
        for i, col in enumerate(sheet["columns"], start=1):
            width = widths[i - 1] if i - 1 < len(widths) else max(12, min(40, len(str(col)) + 4))
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        if sheet.get("notes"):
            ws.cell(row=ws.max_row + 2, column=1, value=str(sheet["notes"])).font = Font(italic=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _latin1(text: Any) -> str:
    # Core PDF fonts are Latin-1 only; swap common Unicode punctuation and drop the rest.
    s = str(text)
    for a, b in {"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-", "•": "-", "…": "...", "☐": "[ ]"}.items():
        s = s.replace(a, b)
    return s.encode("latin-1", "ignore").decode("latin-1")


def build_printable(spec: dict, path: Path) -> Path:
    """Spec: {page_size?: "Letter"|"A4", pages: [{title, subtitle?, sections: [{heading, kind, items|lines}]}]}.

    Section kinds: "checklist" (items with boxes), "lines" (N blank writing lines),
    "table" (columns + rows count), "text" (paragraph).
    """
    pdf = FPDF(format=spec.get("page_size", "Letter"))
    pdf.set_auto_page_break(auto=True, margin=15)
    width = pdf.w - pdf.l_margin - pdf.r_margin
    for page in spec["pages"]:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 22)
        pdf.multi_cell(width, 11, _latin1(page["title"]), align="C", new_x="LMARGIN", new_y="NEXT")
        if page.get("subtitle"):
            pdf.set_font("Helvetica", "I", 12)
            pdf.multi_cell(width, 7, _latin1(page["subtitle"]), align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)
        for section in page.get("sections", []):
            if section.get("heading"):
                pdf.set_font("Helvetica", "B", 13)
                pdf.multi_cell(width, 8, _latin1(section["heading"]), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "", 11)
            kind = section.get("kind", "text")
            if kind == "checklist":
                for item in section.get("items", []):
                    x, y = pdf.get_x(), pdf.get_y()
                    pdf.rect(x, y + 1.5, 4, 4)
                    pdf.set_x(x + 7)
                    pdf.multi_cell(width - 7, 7, _latin1(item), new_x="LMARGIN", new_y="NEXT")
            elif kind == "lines":
                for _ in range(int(section.get("lines", 5))):
                    pdf.ln(8)
                    y = pdf.get_y()
                    pdf.line(pdf.l_margin, y, pdf.l_margin + width, y)
                pdf.ln(2)
            elif kind == "table":
                cols = section.get("columns", [])
                if cols:
                    w = width / len(cols)
                    pdf.set_font("Helvetica", "B", 10)
                    for c in cols:
                        pdf.cell(w, 8, _latin1(c), border=1, align="C")
                    pdf.ln()
                    pdf.set_font("Helvetica", "", 10)
                    for _ in range(int(section.get("rows", 8))):
                        for _c in cols:
                            pdf.cell(w, 8, "", border=1)
                        pdf.ln()
            else:
                pdf.multi_cell(width, 6, _latin1(section.get("text", "")), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(3)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    return path


def build_product(spec: dict, folder: Path) -> Path | None:
    if spec["type"] == "spreadsheet":
        return build_spreadsheet(spec, folder / "product.xlsx")
    if spec["type"] == "printable":
        return build_printable(spec, folder / "product.pdf")
    return None
