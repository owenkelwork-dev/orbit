import json
import shutil
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from orbit import rules
from orbit.agents import Context, forge, scout
from orbit.config import load_config
from orbit.coordinator import default_plan, fix_order, run
from orbit.llm import Chat, Claude, Usage, _echo_content, ask_json, extract_json
from orbit.state import State, read_income

REPO = Path(__file__).resolve().parent.parent
TODAY = date(2026, 10, 5)

GOOD_TAGS = [
    "budget planner", "budget spreadsheet", "monthly budget", "expense tracker", "savings tracker",
    "debt payoff", "finance planner", "money tracker", "bill tracker", "paycheck budget",
    "young adult budget", "digital download", "zero based budget",
]
LISTING = {
    "title": "Monthly Budget Spreadsheet for Young Adults, Simple Expense and Savings Tracker",
    "tags": GOOD_TAGS,
    "description": "A simple monthly budget spreadsheet.",
    "price_usd": 6.5,
    "photo_plan": ["Mockup on laptop", "What's included grid"],
}
SHEET = {"type": "spreadsheet", "sheets": [
    {"name": "Instructions", "columns": ["How to use"], "rows": [["Fill in the Budget tab."]]},
    {"name": "Budget", "columns": ["Category", "Planned", "Actual", "Difference"],
     "rows": [["Rent", 900, 950, "=C2-B2"], ["Food", 300, 280, "=C3-B3"]]},
]}
SCOUT_REPLY = {
    "candidates": [{"name": n, "product_type": "spreadsheet", "score": 7, "note": "x"} for n in ("A", "B", "C")],
    "pick": {"name": "Paycheck Budget Spreadsheet", "product_type": "spreadsheet", "target_buyer": "new grads",
             "why": "My judgment: underserved.", "evidence": [{"claim": "Buyers ask for it", "source_url": "https://example.com/a"}]},
}


def fenced(obj):
    return "Here you go.\n```json\n" + json.dumps(obj) + "\n```"


class FakeChat:
    def __init__(self, replies):
        self.replies = replies
        self.sent = []

    def send(self, text):
        self.sent.append(text)
        return self.replies.pop(0)


class FakeLLM:
    """Hands out scripted replies in order; records which system prompts were used."""

    def __init__(self, replies):
        self.replies = list(replies)
        self.systems = []
        self.usage = Usage(input_tokens=1000, output_tokens=500, web_searches=2, usd=0.05)

    def chat(self, system, web_search=False):
        self.systems.append(system)
        return FakeChat(self.replies)


@pytest.fixture
def repo(tmp_path):
    for name in ("orbit.toml", "prompts", "data"):
        src = REPO / name
        (shutil.copytree if src.is_dir() else shutil.copy)(src, tmp_path / name)
    return load_config(tmp_path)


# ------------------------------------------------------------------ rules

def test_listing_rules_accept_valid_listing():
    assert rules.check_etsy_listing(LISTING) == []


def test_listing_rules_catch_title_and_tag_problems():
    bad = {**LISTING, "title": " ".join(["word"] * 15), "tags": GOOD_TAGS[:12] + ["this tag is way too long"]}
    errors = rules.check_etsy_listing(bad)
    assert any("15 words" in e for e in errors)
    assert any("way too long" in e for e in errors)
    assert rules.check_etsy_listing({**LISTING, "tags": GOOD_TAGS[:12]})  # 12 tags
    assert rules.check_etsy_listing({**LISTING, "tags": GOOD_TAGS[:12] + ["budget planner"]})  # repeat


def test_tag_of_exactly_20_chars_is_allowed():
    tag = "a" * 20
    assert rules.check_etsy_listing({**LISTING, "tags": GOOD_TAGS[:12] + [tag]}) == []


def test_trademark_check():
    assert rules.find_trademarks("Cute Disney Planner") == ["disney"]
    assert rules.find_trademarks("Budget planner for the whole office") == []
    assert rules.check_trademarks({"a": ["Harry Potter themed"]}, "x")


def test_affiliate_post_rules():
    ok = {"markdown": "# Desks\n\n*Disclosure: this post has affiliate links; I may earn a commission.*\n\nBuy [AFFILIATE LINK]"}
    assert rules.check_affiliate_post(ok) == []
    errors = rules.check_affiliate_post({"markdown": "# Desks\n\nBuy https://amzn.to/x"})
    assert len(errors) == 3


# ------------------------------------------------------------------ llm helpers

def test_extract_json_prefers_last_fenced_block():
    text = 'draft ```json\n{"a": 1}\n``` final:\n```json\n{"a": 2}\n```'
    assert extract_json(text) == {"a": 2}
    assert extract_json('sure: {"b": 3} thanks') == {"b": 3}


def test_ask_json_retries_with_errors_then_succeeds():
    chat = FakeChat([fenced({"n": 1}), fenced({"n": 2})])
    data = ask_json(chat, "go", lambda d: [] if d["n"] == 2 else ["n must be 2"])
    assert data == {"n": 2}
    assert "n must be 2" in chat.sent[1]


def test_echo_content_strips_declined_partial_before_fallback():
    blocks = [{"type": "thinking", "thinking": ""}, {"type": "text", "text": "par"},
              {"type": "fallback"}, {"type": "text", "text": "done"}]
    assert [b["type"] for b in _echo_content(blocks)] == ["text", "fallback", "text"]
    assert _echo_content(blocks[1:2]) == blocks[1:2]


def test_chat_resumes_on_pause_turn_and_tracks_usage():
    def resp(stop, text):
        usage = SimpleNamespace(input_tokens=100, output_tokens=10, cache_read_input_tokens=0,
                                cache_creation_input_tokens=0, server_tool_use=SimpleNamespace(web_search_requests=1))
        block = SimpleNamespace(to_dict=lambda: {"type": "text", "text": text})
        return SimpleNamespace(content=[block], stop_reason=stop, model="claude-opus-5-5", usage=usage)

    queue = [resp("pause_turn", "a"), resp("end_turn", "b")]

    class Stream:
        def __init__(self, **kw):
            self.kw = kw

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def get_final_message(self):
            return queue.pop(0)

    calls = []
    client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(stream=lambda **kw: calls.append(kw) or Stream(**kw))))
    llm = Claude("claude-opus-5-5", "medium", 3, client=client)
    out = llm.chat("sys", web_search=True).send("hi")
    assert out == "a\nb"
    assert len(calls) == 2
    assert calls[0]["fallbacks"] == "default" and calls[0]["tools"][0]["name"] == "web_search"
    assert llm.usage.web_searches == 2 and llm.usage.usd > 0


# ------------------------------------------------------------------ state

def test_read_income(tmp_path):
    p = tmp_path / "income.csv"
    p.write_text("date,channel,amount,note\n2026-10-01,Etsy,12.50,first sale\n2026-01-01,fiverr,\"$1,000.00\",x\nbad,etsy,3,\n")
    s = read_income(p, TODAY)
    assert s.total == 1012.5 and s.by_channel == {"etsy": 12.5, "fiverr": 1000.0}
    assert s.last_30_days == 12.5
    assert len(s.bad_rows) == 1


def test_fix_order_runs_scout_before_forge(repo):
    state = State(repo)
    assert fix_order(["forge", "quill"], state, 2) == ["scout", "forge"]
    state.ideas.append({"name": "x", "forged": False})
    assert fix_order(["forge", "quill"], state, 2) == ["forge", "quill"]


def test_default_plan_starts_with_scout(repo):
    assert default_plan(State(repo), TODAY, 2).agents == ["scout", "pixel"]


# ------------------------------------------------------------------ agents

def test_scout_then_forge_builds_spreadsheet(repo):
    llm = FakeLLM([fenced(SCOUT_REPLY), fenced({"listing": LISTING, "product": SHEET})])
    ctx = Context(repo, llm, State(repo), TODAY)
    scout(ctx)
    assert ctx.state.pending_ideas()[0]["name"] == "Paycheck Budget Spreadsheet"
    result = forge(ctx)
    assert not ctx.state.pending_ideas()
    xlsx = next(f for f in result.files if f.suffix == ".xlsx")
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Instructions", "Budget"]
    assert wb["Budget"]["D2"].value == "=C2-B2"
    assert "house rules" in llm.systems[0]


def test_forge_builds_printable_pdf(repo):
    state = State(repo)
    state.ideas.append({"name": "Bill Checklist", "product_type": "printable", "forged": False})
    spec = {"type": "printable", "pages": [{"title": "Bills — Monthly", "sections": [
        {"heading": "Due", "kind": "checklist", "items": ["Rent"]},
        {"heading": "Notes", "kind": "lines", "lines": 3},
        {"heading": "Log", "kind": "table", "columns": ["Date", "Amount"], "rows": 3},
        {"heading": "How", "kind": "text", "text": "Print it."}]}]}
    result = forge(Context(repo, FakeLLM([fenced({"listing": LISTING, "product": spec})]), state, TODAY))
    pdf = next(f for f in result.files if f.suffix == ".pdf")
    assert pdf.read_bytes().startswith(b"%PDF")


# ------------------------------------------------------------------ end to end

def test_run_writes_report_and_state(repo):
    plan = {"agents": ["forge"], "reasoning": "Start with Etsy.", "notes_for_owner": ["Log sales."]}
    llm = FakeLLM([fenced(plan), fenced(SCOUT_REPLY), fenced({"listing": LISTING, "product": SHEET})])
    report = run(repo, lambda: llm, today=TODAY)
    text = report.read_text()
    assert report.name == "2026-10-05.md"
    assert "SCOUT" in text and "FORGE" in text and "Log sales." in text
    assert "Nothing was posted, sent" in text
    state = json.loads((repo.data_dir / "state.json").read_text())
    assert state["last_run"] == {"scout": "2026-10-05", "forge": "2026-10-05"}
    assert (repo.data_dir / "spend.csv").read_text().count("\n") == 2


def test_run_survives_agent_failure_and_bad_plan(repo):
    llm = FakeLLM(["no json here"] * 3 + ["also not json"] * 3 + [fenced({"title": "I will", "tags": [], "packages": []})] * 3)
    text = run(repo, lambda: llm, today=TODAY).read_text()
    assert "fallback plan" in text
    assert "(failed)" in text


def test_run_skips_when_over_budget(repo):
    (repo.data_dir / "spend.csv").write_text("date,model,input_tokens,output_tokens,web_searches,usd\n2026-10-01,m,1,1,0,99\n")
    text = run(repo, lambda: pytest.fail("should not create a client"), today=TODAY).read_text()
    assert "Run skipped" in text
